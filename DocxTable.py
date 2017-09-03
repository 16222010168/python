#coding:utf-8
from docx import Document
from xlwt import Workbook

document = Document('openSSH.docx')
table = document.tables
book = Workbook()
sheet1 = book.add_sheet('Sheet1')
TbCount = len(table)

for i in range(0,TbCount):
    tmp = []
    if table[i].cell(0,0).text == u'漏洞编号':
        pass
    else:
        ColCount = len(table[i].columns)
        RowCount = len(table[i].rows)
        for r in range(0, RowCount):
            for c in range(0, ColCount):
                tmp.append(table[i].cell(r,c).text)

    for ExlRow in range(0,len(tmp)):
        sheet1.write(ExlRow,i,tmp[ExlRow])

book.save('table.xls')









